import csv
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


BRAND_NAME = "Get Nailed Nail Bar and Spa"
DEEP_ROSE = "8B4F5C"


def _safe_cell(value):
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def _filename(report, extension):
    return f"{report.key}-{report.start_date}-{report.end_date}.{extension}"


def export_csv(report):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{_filename(report, "csv")}"'
    response.write("\ufeff")
    writer = csv.writer(response)
    writer.writerow(report.columns)
    for row in report.rows:
        writer.writerow([_safe_cell(row[column]) for column in report.columns])
    return response


def export_xlsx(report):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report.title[:31]
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(report.columns))
    title = sheet.cell(1, 1, f"{BRAND_NAME} - {report.title}")
    title.font = Font(size=16, bold=True, color=DEEP_ROSE)
    title.alignment = Alignment(horizontal="center")
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(report.columns))
    sheet.cell(2, 1, f"{report.start_date} to {report.end_date}").alignment = Alignment(
        horizontal="center"
    )
    for column_index, column in enumerate(report.columns, 1):
        cell = sheet.cell(4, column_index, column)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=DEEP_ROSE)
    for row_index, row in enumerate(report.rows, 5):
        for column_index, column in enumerate(report.columns, 1):
            value = row[column]
            if isinstance(value, str):
                value = _safe_cell(value)
            sheet.cell(row_index, column_index, value)
    for index, column in enumerate(report.columns, 1):
        values = [str(column), *(str(row[column]) for row in report.rows)]
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(len(value) for value in values) + 3, 40
        )
    sheet.freeze_panes = "A5"
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{_filename(report, "xlsx")}"'
    return response


def export_pdf(report):
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
        title=report.title,
        author=BRAND_NAME,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "BrandTitle",
        parent=styles["Title"],
        textColor=colors.HexColor(f"#{DEEP_ROSE}"),
        alignment=TA_CENTER,
    )
    story = []
    logo_path = Path(settings.BASE_DIR) / "static" / "images" / "get-nailed-logo.png"
    if logo_path.exists():
        story.append(Image(str(logo_path), width=23 * mm, height=23 * mm))
    story.extend(
        [
            Paragraph(BRAND_NAME, title_style),
            Paragraph(report.title, styles["Heading2"]),
            Paragraph(f"{report.start_date} to {report.end_date}", styles["Normal"]),
            Spacer(1, 5 * mm),
        ]
    )
    data = [list(report.columns)]
    data.extend(
        [[_safe_cell(row[column]) for column in report.columns] for row in report.rows]
    )
    table = Table(data, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{DEEP_ROSE}")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#E6D9DA")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FFF7F7")]),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(table)
    document.build(story)
    response = HttpResponse(output.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{_filename(report, "pdf")}"'
    return response
